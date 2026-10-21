#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_equipment_detail.py — тест модуля Equipment_Detail_Export.

Проверяет на примере examples/Dorado 6000 V3_BBU.txt:
1) распаршено ровно 6 записей BBU;
2) статический RemainLife из лога = 780;
3) Actual_RemainLife ОТЛИЧАЕТСЯ от 780 (лог старый — прошло > 780 дней,
   поэтому ожидается "Expired");
4) Excel-файл создан, содержит 6 строк данных и стилизацию заголовков;
5) unit-тесты расчёта calculate_actual_remainlife (positive / expired / None).

Запуск:  python test_equipment_detail.py
       (или pytest test_equipment_detail.py)
"""

import os
import sys
import tempfile
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from Equipment_Detail_Export import (
    export_equipment_detail,
    parse_log_date,
    calculate_actual_remainlife,
    STATUS_EXPIRED,
)

LOG_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "examples", "Dorado 6000 V3_BBU.txt",
)

STATIC_REMAIN_LIFE_IN_LOG = 780


# ==============================================================================
# 1. Интеграционный тест: экспорт + содержимое Excel
# ==============================================================================

def test_export_equipment_detail():
    assert os.path.isfile(LOG_FILE), f"Test log not found: {LOG_FILE}"

    out_path = os.path.join(tempfile.mkdtemp(), "BBU_Detail.xlsx")
    result = export_equipment_detail(LOG_FILE, out_path)

    records = result["records"]

    # --- Ровно 6 строк данных ---
    assert len(records) == 6, (
        f"Expected 6 BBU records, got {len(records)}"
    )

    # --- Статический RemainLife в логе = 780 ---
    statics = [r["remain_life_static"] for r in records]
    assert all(s == STATIC_REMAIN_LIFE_IN_LOG for s in statics), (
        f"Static RemainLife mismatch: {statics}"
    )

    # --- Дата лога распознана ---
    log_date = result["log_date"]
    assert log_date == datetime(2023, 5, 10, 9, 30, 0), (
        f"Log date parsing failed: {log_date}"
    )

    # --- Лог старый: прошло больше, чем 780 дней ---
    days_since = result["days_since_log"]
    assert days_since > STATIC_REMAIN_LIFE_IN_LOG, (
        f"Test log expected to be old (>780 days), days_since={days_since}"
    )

    # --- Actual_RemainLife отличается от статического 780 ---
    actual_values = [
        calculate_actual_remainlife(r["remain_life_static"], log_date)
        for r in records
    ]
    for actual in actual_values:
        assert actual != STATIC_REMAIN_LIFE_IN_LOG, (
            "Actual RemainLife must differ from static 780!"
        )
        assert actual == STATUS_EXPIRED, (
            f"Old log with 780-day life should be Expired, got {actual!r}"
        )

    # --- Обязательные поля извлечены ---
    for r in records:
        assert r["bbu_id"], "Missing BBU ID"
        assert r["enclosure_id"], "Missing Enclosure ID"
        assert r["health_status"], "Missing Health Status"
        assert r["running_status"], "Missing Running Status"
        assert r["voltage"], "Missing Voltage(V)"
        assert r["barcode"], "Missing BarCode"
        assert r["description"], "Missing Description"
        assert r["delivered_on"], "Missing Delivered On"
        assert r["manufactured_on"], "Missing Manufactured="

    # --- Excel создан и содержит 6 строк данных ---
    assert os.path.isfile(out_path), "Excel file was not created"
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    header_row = 3
    data_rows = []
    row = header_row + 1
    while ws.cell(row=row, column=2).value:
        data_rows.append(row)
        row += 1

    assert len(data_rows) == 6, (
        f"Expected 6 data rows in Excel, got {len(data_rows)}"
    )

    # Колонка "Actual RemainLife (days)" — последняя
    actual_col = len(result.get("columns", [])) or 14
    headers = [ws.cell(row=header_row, column=c).value
               for c in range(1, 15)]
    assert headers[-1] == "Actual RemainLife (days)", headers
    actual_col = 14

    for dr in data_rows:
        cell_val = ws.cell(row=dr, column=actual_col).value
        assert cell_val != STATIC_REMAIN_LIFE_IN_LOG, (
            f"Excel row {dr}: Actual RemainLife equals static 780!"
        )
        assert cell_val == STATUS_EXPIRED, (
            f"Excel row {dr}: expected 'Expired', got {cell_val!r}"
        )

    # --- Стилизация: жирный заголовок, границы ячеек ---
    h_cell = ws.cell(row=header_row, column=1)
    assert h_cell.font.bold, "Header font must be bold"
    assert h_cell.border.left.style, "Header cell must have borders"
    d_cell = ws.cell(row=data_rows[0], column=1)
    assert d_cell.border.left.style, "Data cell must have borders"

    print(f"[OK] Export test passed: {len(records)} records, "
          f"days_since_log={days_since}, Actual RemainLife='{STATUS_EXPIRED}'")
    return True


# ==============================================================================
# 2. Unit-тесты расчёта RemainLife
# ==============================================================================

def test_calculate_actual_remainlife():
    log_dt = datetime(2024, 1, 1, 0, 0, 0)

    # Прошло 100 дней → 780 - 100 = 680
    v = calculate_actual_remainlife(780, log_dt, now=datetime(2024, 4, 10))
    assert v == 680, v

    # Ровно в день лога → без изменений
    v = calculate_actual_remainlife(780, log_dt, now=log_dt)
    assert v == 780, v

    # Прошло больше срока → Expired
    v = calculate_actual_remainlife(780, log_dt, now=datetime(2027, 1, 1))
    assert v == STATUS_EXPIRED, v

    # Нет данных → None
    assert calculate_actual_remainlife(None, log_dt) is None
    assert calculate_actual_remainlife(780, None) is None

    # Будущая дата лога (часовой пояс и т.п.) → не отрицательный сдвиг
    v = calculate_actual_remainlife(780, datetime(2030, 1, 1), now=log_dt)
    assert v == 780, v

    print("[OK] calculate_actual_remainlife unit tests passed")
    return True


def test_parse_log_date():
    content = "PROFILE FOR STORAGE ARRAY: MyArray (2023-5-10 09:30:00)\n..."
    assert parse_log_date(content) == datetime(2023, 5, 10, 9, 30, 0)
    assert parse_log_date("no header here") is None
    print("[OK] parse_log_date unit tests passed")
    return True


# ==============================================================================
# RUNNER
# ==============================================================================

def main():
    tests = [
        test_parse_log_date,
        test_calculate_actual_remainlife,
        test_export_equipment_detail,
    ]
    failed = 0
    for t in tests:
        try:
            t()
        except AssertionError as e:
            failed += 1
            print(f"[FAIL] {t.__name__}: {e}")
        except Exception as e:
            failed += 1
            print(f"[ERROR] {t.__name__}: {type(e).__name__}: {e}")

    total = len(tests)
    print(f"\nResult: {total - failed} OK, {failed} FAIL")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())

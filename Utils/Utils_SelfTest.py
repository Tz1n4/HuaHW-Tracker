#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Utils_SelfTest.py — self-test для HuaHW-Tracker.
"""

import os
import re
import sys


# ==============================================================================
# ВСПОМОГАТЕЛЬНЫЕ
# ==============================================================================

def _extract_placeholders(template):
    """Извлекает имена {placeholder} из строки-шаблона."""
    return set(re.findall(r'\{(\w+)\}', template))


# ==============================================================================
# ТОЧКА ВХОДА
# ==============================================================================

def run_checks(test_file=None):
    """Run all self-test checks. Returns list of result dicts."""
    results = []

    def add(label, ok, extra=""):
        results.append({"label": label, "ok": bool(ok), "extra": extra})

    script_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    # --------------------------------------------------------------------------
    # ОКРУЖЕНИЕ
    # --------------------------------------------------------------------------

    add("Python interpreter", True, f"v{sys.version.split()[0]}")
    add("Script folder in sys.path", script_dir in sys.path, script_dir)

    # Обновленный список файлов с учетом новой структуры папок
    required_files = [
        "Utils/Utils_Constants.py", "Utils/Utils_Regex_Helpers.py", "Utils/Utils_System.py",
        "Utils/Utils_Locale.py", "Utils/Utils_SelfTest.py",
        "Core/Core_System_Info.py", "Core/Core_Base_Parser.py", "Core/Core_Export_Excel.py",
        "Parsers/Parser_V3.py", "Parsers/Parser_V5.py", "Parsers/Parser_V6.py",
        "Parsers/Parsers_Details/Parser_BBU_Details.py",
        "GUI/Gui_Workers.py", "GUI/Gui_App.py", "GUI/Gui_Main.py", "GUI/Gui_Settings.py",
        "GUI/Gui_Dialogs/Gui_Dialog_Detailed_Analysis.py",
        "Main.py"
    ]
    missing = [f for f in required_files if not os.path.exists(os.path.join(script_dir, f))]
    add(
        f"Script files ({len(required_files)} files)",
        not missing,
        "all present" if not missing else f"missing: {missing}",
    )

    # Проверка импорта парсеров
    try:
        from Parsers.Parser_V3 import V3Parser          # noqa: F401
        from Parsers.Parser_V5 import V5Parser          # noqa: F401
        from Parsers.Parser_V6 import V6Parser          # noqa: F401
        add("Parser modules import (V3/V5/V6)", True, "3 parsers")
    except ImportError as e:
        add("Parser modules import (V3/V5/V6)", False, str(e))

    # Проверка импорта парсера деталей BBU
    try:
        from Parsers.Parsers_Details.Parser_BBU_Details import parse_bbu_details  # noqa: F401
        add("BBU Details parser import", True, "parse_bbu_details")
    except ImportError as e:
        add("BBU Details parser import", False, str(e))

    template = os.path.join(script_dir, "Form.xlsx")
    add(
        "Template Form.xlsx exists",
        os.path.exists(template),
        template if os.path.exists(template) else "NOT FOUND",
    )

    try:
        import openpyxl
        add("openpyxl", True, f"v{openpyxl.__version__}")
    except ImportError as e:
        add("openpyxl", False, str(e))

    try:
        from PyQt6.QtCore import QT_VERSION_STR
        add("PyQt6 (Qt)", True, f"Qt {QT_VERSION_STR}")
    except ImportError as e:
        add("PyQt6", False, str(e))

    # --------------------------------------------------------------------------
    # ИМПОРТЫ ЯДРА
    # --------------------------------------------------------------------------

    try:
        from Core.Core_Base_Parser import HuaweiStorageParser    # noqa: F401
        add("Parser module import", True, "HuaweiStorageParser")
    except ImportError as e:
        add("Parser module import", False, str(e))

    try:
        from Core.Core_Export_Excel import ExcelExporter              # noqa: F401
        add("Exporter module import", True, "ExcelExporter")
    except ImportError as e:
        add("Exporter module import", False, str(e))

    # --------------------------------------------------------------------------
    # ЛОКАЛИЗАЦИЯ
    # --------------------------------------------------------------------------

    try:
        from Utils.Utils_Locale import STRINGS
        en_keys = set(STRINGS["en"].keys())
        ru_keys = set(STRINGS["ru"].keys())

        if en_keys == ru_keys:
            add("Locale keys (EN/RU match)", True, f"{len(en_keys)} keys")
        else:
            only_en = sorted(en_keys - ru_keys)
            only_ru = sorted(ru_keys - en_keys)
            detail = []
            if only_en:
                detail.append(f"only EN: {only_en[:3]}")
            if only_ru:
                detail.append(f"only RU: {only_ru[:3]}")
            add("Locale keys (EN/RU match)", False, "; ".join(detail))
    except Exception as e:
        add("Locale keys (EN/RU match)", False, str(e))

    try:
        from Utils.Utils_Locale import STRINGS
        bad = []
        for lang, strings in STRINGS.items():
            for key, tmpl in strings.items():
                if not isinstance(tmpl, str) or "{" not in tmpl:
                    continue
                placeholders = _extract_placeholders(tmpl)
                probe = {name: "x" for name in placeholders}
                try:
                    tmpl.format(**probe)
                except (KeyError, IndexError, ValueError) as e:
                    bad.append(f"{lang}/{key}: {e}")

        if not bad:
            add("Locale templates well-formed", True, "all OK")
        else:
            add("Locale templates well-formed", False, "; ".join(bad[:3]))
    except Exception as e:
        add("Locale templates well-formed", False, str(e))

    # --------------------------------------------------------------------------
    # REGEX-ХЕЛПЕРЫ (smoke-тесты)
    # --------------------------------------------------------------------------

    try:
        from Utils.Utils_Regex_Helpers import (
            extract_label_fields,
            extract_part_number_from_serial,
            sanitize_for_excel,
        )

        sn = "2102353BBF9WM5800015"
        pn = extract_part_number_from_serial(sn)
        if pn == "02353BBF":
            add("Regex: part number from serial", True, f"{sn} → {pn}")
        else:
            add("Regex: part number from serial", False, f"got {pn!r}, expected '02353BBF'")

        block = "Description=Test Module\nBarCode=SN123\nItem=02351ABC"
        lbl = extract_label_fields(block)
        if lbl["description"] == "Test Module" and lbl["item"] == "02351ABC":
            add("Regex: elabel parse", True, "Description/Item OK")
        else:
            add("Regex: elabel parse", False, str(lbl))

        dirty = "Test\x00Module &amp; Co"
        clean = sanitize_for_excel(dirty)
        if clean == "TestModule & Co":
            add("Regex: sanitize for excel", True, repr(clean))
        else:
            add("Regex: sanitize for excel", False, repr(clean))

    except Exception as e:
        add("Regex helpers", False, str(e))

    # --------------------------------------------------------------------------
    # ЛОГИКА ПАРСЕРА
    # --------------------------------------------------------------------------

    try:
        from Core.Core_System_Info import detect_version

        cases = [
            ("OceanStor Dorado 3000 V6", "", "", "V6"),
            ("OceanStor 5500 V5", "", "", "V5"),
            ("OceanStor 5300 V3", "", "", "V3"),
            ("", "V600R005C27", "", "V6"),
            ("", "V500R007C20", "", "V5"),
            ("", "", "Control Board(0A)", "V6"),
            ("", "", "Controller ID: 0A\nsfp info:", "V3"),
            ("", "", "Controller ID: 0A", "V5"),
            ("", "", "nothing relevant", "Unknown"),
        ]

        failed = []
        for model, fw, content, expected in cases:
            got = detect_version(model=model, fw=fw, content=content)
            if got != expected:
                hint = model or fw or content[:30].replace("\n", "\\n")
                failed.append(f"{hint!r} → {got} (want {expected})")

        if not failed:
            add("Version detect (9 cases)", True, "all pass")
        else:
            add("Version detect (9 cases)", False, "; ".join(failed[:3]))
    except Exception as e:
        add("Version detect (9 cases)", False, str(e))

    # --------------------------------------------------------------------------
    # ШАБЛОН EXCEL
    # --------------------------------------------------------------------------

    if os.path.exists(template):
        try:
            import openpyxl
            from Utils.Utils_Constants import SHEET_ZIP
            wb = openpyxl.load_workbook(template)
            if SHEET_ZIP in wb.sheetnames:
                add("Template opens (sheet 'ЗИП')", True, SHEET_ZIP)
            else:
                add(
                    "Template opens (sheet 'ЗИП')", False,
                    f"missing sheet '{SHEET_ZIP}', have: {wb.sheetnames}",
                )
        except Exception as e:
            add("Template opens (sheet 'ЗИП')", False, str(e))
    else:
        add("Template opens (sheet 'ЗИП')", False, "skipped (file missing)")

    # --------------------------------------------------------------------------
    # РЕАЛЬНЫЙ РАЗБОР (опционально)
    # --------------------------------------------------------------------------

    if test_file and os.path.isfile(test_file) and test_file.lower().endswith(".txt"):
        try:
            from Core.Core_Base_Parser import HuaweiStorageParser
            parser = HuaweiStorageParser()
            sys_info, comps = parser.parse_file(test_file)
            add(
                "Test parse (selected file)", True,
                f"{sys_info['name']} — {len(comps)} components",
            )
        except Exception as e:
            add("Test parse (selected file)", False, str(e))
    else:
        add("Test parse (no file selected)", True, "skipped")

    return results
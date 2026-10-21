#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Core_Export_Excel_Details.py — генерация Excel-отчёта по детальному анализу.
"""

import os
import shutil
from datetime import datetime
from copy import copy

import openpyxl

from Utils.Utils_Constants import (
    DETAILS_TEMPLATE_PATH,
    TIMESTAMP_FORMAT,
)
from Utils.Utils_Regex_Helpers import sanitize_for_excel


class DetailsExcelExporter:
    """Формирует Excel-отчёт по детальному анализу компонентов."""

    def __init__(self, template_path=None):
        self.template_path = template_path or DETAILS_TEMPLATE_PATH
        if not os.path.exists(self.template_path):
            raise FileNotFoundError(f"Template not found: {self.template_path}")

    def export_bbu(self, bbu_data, output_dir, name_hint):
        """Экспорт данных BBU в Excel."""
        timestamp = datetime.now().strftime(TIMESTAMP_FORMAT)
        filename = f"Детальный_анализ_{name_hint}_{timestamp}.xlsx"
        output_path = os.path.join(output_dir, filename)

        shutil.copyfile(self.template_path, output_path)
        wb = openpyxl.load_workbook(output_path)
        
        if "BBU" not in wb.sheetnames:
            raise ValueError("Шаблон Form_Details.xlsx не содержит лист 'BBU'")
            
        ws = wb["BBU"]
        current_row = 2  # Начинаем заполнять со 2-й строки

        for i, bbu in enumerate(bbu_data, start=1):
            # Копируем зеленый стиль со второй строки на все последующие
            if current_row > 2:
                self._copy_style(ws, src_row=2, dst_row=current_row)

            # Заполняем данные (A - N)
            self._set(ws, current_row, 1, i)
            self._set(ws, current_row, 2, bbu.get("sys_name", ""))
            self._set(ws, current_row, 3, bbu.get("sys_model", ""))
            self._set(ws, current_row, 4, bbu.get("id", ""))
            self._set(ws, current_row, 5, bbu.get("health", ""))
            self._set(ws, current_row, 6, bbu.get("running", ""))
            self._set(ws, current_row, 7, bbu.get("voltage", ""))
            
            self._set(ws, current_row, 8, bbu.get("remain_life_log", "N/A"))
            self._set(ws, current_row, 9, bbu.get("real_remain_days", "N/A"))
            self._set(ws, current_row, 10, bbu.get("end_of_life", ""))
            self._set(ws, current_row, 11, bbu.get("delivered_on", ""))
            # Форматируем возраст в текст (как в интерфейсе)
            age_days = bbu.get("age_days")
            if isinstance(age_days, int):
                age_years = age_days // 365
                age_rem_days = age_days % 365
                if age_years > 0:
                    age_str = f"{age_years} г. {age_rem_days} дн."
                else:
                    age_str = f"{age_rem_days} дн."
                self._set(ws, current_row, 12, age_str)
            else:
                self._set(ws, current_row, 12, "N/A")
            self._set(ws, current_row, 13, bbu.get("item", ""))
            self._set(ws, current_row, 14, bbu.get("barcode", ""))

            current_row += 1

        wb.save(output_path)
        return output_path

    def _set(self, ws, row, col, value):
        """Запись значения с санитизацией."""
        ws.cell(row=row, column=col).value = sanitize_for_excel(value)

    def _copy_style(self, ws, src_row, dst_row):
        """Копирует стиль всех ячеек из src_row в dst_row."""
        total_cols = ws.max_column
        for col in range(1, total_cols + 1):
            src = ws.cell(row=src_row, column=col)
            dst = ws.cell(row=dst_row, column=col)
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)